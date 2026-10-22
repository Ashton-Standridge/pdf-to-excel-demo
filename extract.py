"""Extract purchase-agreement fields from a PDF into an Excel workbook with a key-date timeline."""
import argparse, re, sys
from datetime import date, datetime, timedelta
from pathlib import Path
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

LABELS = ["Buyer", "Seller", "Property Address", "Tax Parcel ID", "Purchase Price", "Earnest Money",
          "Earnest Money Holder", "Seller Contribution at Closing", "Financing Type", "Due Diligence Period",
          "Earnest Money Due", "Financing Contingency", "Appraisal Contingency", "Closing", "Possession"]
DEADLINES = ["Earnest Money Due", "Due Diligence Period", "Appraisal Contingency", "Financing Contingency", "Closing"]

def read_text(pdf: Path, password: str | None = None) -> str:
    try:
        with pdfplumber.open(pdf, password=password or "") as doc:
            return "\n".join(p.extract_text() or "" for p in doc.pages)
    except Exception as e:
        if "password" in str(e).lower() or type(e).__name__ in ("PDFPasswordIncorrect", "PdfminerException"):
            raise SystemExit(f"Could not open {pdf.name}: it is password-protected or the password is wrong. Use --password.")
        raise

def parse_fields(text: str) -> dict:
    out = {}
    for label in LABELS:
        m = re.search(rf"^{re.escape(label)}:\s*(.+)$", text, re.M)
        out[label] = m.group(1).strip() if m else ""
    return out

def roll_weekend(d: date) -> date:
    return d + timedelta(days={5: 2, 6: 1}.get(d.weekday(), 0))

def timeline(fields: dict, accepted: date) -> list[tuple]:
    rows = [("Binding Agreement Date", 0, accepted, accepted, "")]
    for k in DEADLINES:
        m = re.search(r"(\d+)", fields.get(k, ""))
        if not m: continue
        n = int(m.group(1)); raw = accepted + timedelta(days=n); adj = roll_weekend(raw)
        rows.append((k, n, raw, adj, "Rolled from weekend to Monday" if adj != raw else ""))
    return rows

def money(s):
    try: return float(s.replace("$", "").replace(",", ""))
    except ValueError: return s

def write_xlsx(fields, rows, out: Path, accepted: date):
    wb = Workbook(); hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F4E78")
    ws = wb.active; ws.title = "Fields"; ws.append(["Field", "Value"])
    for k, v in fields.items():
        ws.append([k, money(v) if v.startswith("$") else v])
        if v.startswith("$"): ws.cell(ws.max_row, 2).number_format = '"$"#,##0.00'
    t = wb.create_sheet("Timeline")
    t["A1"], t["B1"] = "Accepted date (edit me):", accepted; t["B1"].number_format = "mm/dd/yyyy"
    t["B1"].fill = PatternFill("solid", fgColor="FFF2CC"); t["A1"].font = Font(bold=True)
    t.append([]); t.append(["Milestone", "Days", "Calendar date", "Due date (weekend -> Monday)", "Day", "Note"])
    for i, (name, n, raw, adj, note) in enumerate(rows, start=4):
        t.append([name, n, f"=$B$1+B{i}", f"=C{i}+IF(WEEKDAY(C{i},2)=6,2,IF(WEEKDAY(C{i},2)=7,1,0))",
                  f'=TEXT(D{i},"dddd")', f'=IF(D{i}<>C{i},"Rolled from weekend to Monday","")'])
        for c in "CD": t[f"{c}{i}"].number_format = "ddd mm/dd/yyyy"
    for sh, hrow in ((ws, 1), (t, 3)):
        for c in sh[hrow]: c.font, c.fill = hdr, fill
        for col, w in zip("ABCDEF", (32, 44, 18, 30, 12, 32)): sh.column_dimensions[col].width = w
    wb.save(out)

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pdf", type=Path); ap.add_argument("-o", "--out", type=Path, default=Path("output.xlsx"))
    ap.add_argument("--accepted", default=date.today().isoformat(), help="Binding agreement date, YYYY-MM-DD")
    ap.add_argument("--password")
    a = ap.parse_args(argv)
    accepted = datetime.strptime(a.accepted, "%Y-%m-%d").date()
    fields = parse_fields(read_text(a.pdf, a.password))
    missing = [k for k, v in fields.items() if not v]
    rows = timeline(fields, accepted); write_xlsx(fields, rows, a.out, accepted)
    print(f"Extracted {len(fields)-len(missing)}/{len(fields)} fields -> {a.out}")
    if missing: print("Missing:", ", ".join(missing))
    print(f"\n{'Milestone':26}{'Days':>5}  {'Due date':16}Note")
    for name, n, raw, adj, note in rows: print(f"{name:26}{n:>5}  {adj:%a %m/%d/%Y}  {note}")
    return fields, rows

if __name__ == "__main__":
    main()
